"""
Lamplight Knowledge Library, Reading Desk & Personal Learning Studio Module
Pillar 4: Watch-Folder Knowledge Library, FTS5 Full-Text Search,
Interactive Reading Desk, Research Notebook, and Spaced-Repetition Learning Studio.
"""

from __future__ import annotations

import datetime
import html
import json
import logging
import math
import os
import re
import uuid
import zipfile
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET

import db

logger = logging.getLogger("lamplight.knowledge")

SUPPORTED_EXTENSIONS = {".pdf", ".epub", ".docx", ".doc", ".md", ".txt", ".json", ".csv", ".html"}


# ==============================================================================
# Database Schema Initialization
# ==============================================================================

def init_knowledge_tables():
    """Ensure all Knowledge Library, Notes, Reading, and Study tables exist."""
    with db.conn() as c:
        # 1. Watched Folders
        c.execute(
            """CREATE TABLE IF NOT EXISTS knowledge_sources(
                id TEXT PRIMARY KEY,
                workspace_id TEXT DEFAULT 'ws_default',
                folder_path TEXT NOT NULL,
                name TEXT NOT NULL,
                recursive INTEGER DEFAULT 1,
                auto_sync INTEGER DEFAULT 1,
                last_scanned_at DATETIME,
                file_count INTEGER DEFAULT 0,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP)"""
        )
        c.execute("CREATE INDEX IF NOT EXISTS idx_k_sources_ws ON knowledge_sources(workspace_id)")

        # 2. Knowledge Documents (Indexed documents from watch folders and uploads)
        c.execute(
            """CREATE TABLE IF NOT EXISTS knowledge_documents(
                id TEXT PRIMARY KEY,
                source_id TEXT,
                workspace_id TEXT DEFAULT 'ws_default',
                filename TEXT NOT NULL,
                filepath TEXT NOT NULL,
                file_type TEXT NOT NULL,
                file_size INTEGER NOT NULL,
                title TEXT,
                char_count INTEGER DEFAULT 0,
                page_count INTEGER,
                summary TEXT DEFAULT '',
                mtime REAL DEFAULT 0,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                updated_at DATETIME DEFAULT CURRENT_TIMESTAMP)"""
        )
        c.execute("CREATE INDEX IF NOT EXISTS idx_k_docs_ws ON knowledge_documents(workspace_id)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_k_docs_source ON knowledge_documents(source_id)")

        # 3. FTS5 Virtual Table for Instant Snippet-Ranked Search
        try:
            c.execute(
                """CREATE VIRTUAL TABLE IF NOT EXISTS knowledge_fts USING fts5(
                    doc_id UNINDEXED,
                    workspace_id UNINDEXED,
                    title,
                    filename,
                    content,
                    tokenize='porter unicode61'
                )"""
            )
        except Exception as e:
            logger.warning("FTS5 table init: %s", e)

        # 4. Research Notebook & Sourced Annotations
        c.execute(
            """CREATE TABLE IF NOT EXISTS research_notes(
                id TEXT PRIMARY KEY,
                workspace_id TEXT DEFAULT 'ws_default',
                doc_id TEXT,
                doc_title TEXT DEFAULT '',
                passage TEXT NOT NULL,
                note TEXT DEFAULT '',
                tags TEXT DEFAULT '',
                citation TEXT DEFAULT '',
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP)"""
        )
        c.execute("CREATE INDEX IF NOT EXISTS idx_res_notes_ws ON research_notes(workspace_id)")

        # 5. Reading Desk Progress & Bookmarks
        c.execute(
            """CREATE TABLE IF NOT EXISTS reading_progress(
                id TEXT PRIMARY KEY,
                doc_id TEXT UNIQUE NOT NULL,
                workspace_id TEXT DEFAULT 'ws_default',
                progress_pct REAL DEFAULT 0.0,
                last_position TEXT DEFAULT '',
                bookmarks_json TEXT DEFAULT '[]',
                last_read_at DATETIME DEFAULT CURRENT_TIMESTAMP)"""
        )

        # 6. Learning Studio & Spaced Repetition Flashcards (SM-2 Algorithm)
        c.execute(
            """CREATE TABLE IF NOT EXISTS study_flashcards(
                id TEXT PRIMARY KEY,
                deck_id TEXT NOT NULL,
                deck_title TEXT NOT NULL,
                workspace_id TEXT DEFAULT 'ws_default',
                front TEXT NOT NULL,
                back TEXT NOT NULL,
                difficulty REAL DEFAULT 2.5,
                interval_days INTEGER DEFAULT 1,
                repetitions INTEGER DEFAULT 0,
                next_review_at DATETIME,
                last_reviewed_at DATETIME,
                created_at DATETIME DEFAULT CURRENT_TIMESTAMP)"""
        )
        c.execute("CREATE INDEX IF NOT EXISTS idx_cards_deck ON study_flashcards(deck_id)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_cards_review ON study_flashcards(next_review_at)")


# ==============================================================================
# Document Parsing & Content Extraction
# ==============================================================================

def strip_html_tags(text: str) -> str:
    """Cleanly strip HTML tags and unescape entities for plain text reading."""
    clean = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", text, flags=re.DOTALL | re.IGNORECASE)
    clean = re.sub(r"<br\s*/?>", "\n", clean, flags=re.IGNORECASE)
    clean = re.sub(r"</p>", "\n\n", clean, flags=re.IGNORECASE)
    clean = re.sub(r"</?(h[1-6]|div|tr|li)[^>]*>", "\n", clean, flags=re.IGNORECASE)
    clean = re.sub(r"<[^>]+>", "", clean)
    clean = html.unescape(clean)
    clean = re.sub(r"\n{3,}", "\n\n", clean)
    return clean.strip()


def parse_epub(filepath: Path) -> dict[str, Any]:
    """
    Parse EPUB file into structured metadata, ordered chapters, and full text.
    Uses Python's standard library zipfile module for zero-dependency reliability.
    """
    chapters = []
    title = filepath.stem
    author = "Unknown Author"

    try:
        with zipfile.ZipFile(filepath, "r") as zf:
            namelist = zf.namelist()
            # 1. Locate container.xml to find the root OPF file
            container_path = "META-INF/container.xml"
            opf_path = None
            if container_path in namelist:
                root = ET.fromstring(zf.read(container_path))
                for rootfile in root.iter("{urn:oasis:names:tc:opendocument:xmlns:container}rootfile"):
                    opf_path = rootfile.attrib.get("full-path")
                    if opf_path:
                        break

            if not opf_path:
                # Fallback: look for any .opf file in the archive
                opf_candidates = [n for n in namelist if n.endswith(".opf")]
                if opf_candidates:
                    opf_path = opf_candidates[0]

            spine_items = []
            manifest = {}
            opf_dir = ""

            if opf_path and opf_path in namelist:
                opf_dir = str(Path(opf_path).parent)
                if opf_dir == ".":
                    opf_dir = ""
                opf_data = zf.read(opf_path)
                opf_root = ET.fromstring(opf_data)

                # Extract title and author
                for elem in opf_root.iter():
                    tag = elem.tag.lower()
                    if tag.endswith("title") and elem.text and title == filepath.stem:
                        title = elem.text.strip()
                    elif tag.endswith("creator") and elem.text:
                        author = elem.text.strip()

                # Extract manifest mapping (id -> href)
                for item in opf_root.iter():
                    if item.tag.lower().endswith("item"):
                        i_id = item.attrib.get("id")
                        href = item.attrib.get("href")
                        if i_id and href:
                            manifest[i_id] = href

                # Extract spine (reading order)
                for itemref in opf_root.iter():
                    if itemref.tag.lower().endswith("itemref"):
                        idref = itemref.attrib.get("idref")
                        if idref and idref in manifest:
                            spine_items.append(manifest[idref])

            # If spine could not be parsed, fallback to sorted xhtml/html files
            if not spine_items:
                spine_items = sorted([
                    n for n in namelist
                    if n.lower().endswith((".xhtml", ".html", ".htm")) and not n.lower().endswith("nav.xhtml")
                ])

            full_text_parts = []
            for idx, href in enumerate(spine_items):
                # Resolve href path relative to opf directory
                full_href = f"{opf_dir}/{href}" if opf_dir else href
                full_href = full_href.replace("//", "/")
                if full_href in namelist:
                    content_bytes = zf.read(full_href)
                    try:
                        raw_html = content_bytes.decode("utf-8")
                    except UnicodeDecodeError:
                        raw_html = content_bytes.decode("latin-1", errors="replace")

                    chapter_text = strip_html_tags(raw_html)
                    if chapter_text.strip():
                        # Determine chapter title if available
                        ch_title_match = re.search(r"<h[1-3][^>]*>(.*?)</h[1-3]>", raw_html, re.IGNORECASE | re.DOTALL)
                        ch_title = strip_html_tags(ch_title_match.group(1)) if ch_title_match else f"Chapter {idx + 1}"
                        if len(ch_title) > 60:
                            ch_title = ch_title[:60] + "…"

                        chapters.append({
                            "index": idx + 1,
                            "title": ch_title,
                            "content": chapter_text,
                            "char_count": len(chapter_text)
                        })
                        full_text_parts.append(f"## {ch_title}\n\n{chapter_text}")

            full_text = "\n\n".join(full_text_parts)
            return {
                "title": title,
                "author": author,
                "chapters": chapters,
                "page_count": max(1, len(chapters)),
                "text_content": full_text,
                "char_count": len(full_text)
            }
    except Exception as e:
        logger.error("Error parsing EPUB %s: %s", filepath, e)
        return {
            "title": filepath.stem,
            "author": "Unknown",
            "chapters": [],
            "page_count": 1,
            "text_content": f"[Error reading EPUB: {e}]",
            "char_count": 0
        }


def parse_document(filepath: Path | str) -> dict[str, Any]:
    """Extract full text, metadata, and chapter/page breakdown from any supported document."""
    path = Path(filepath)
    if not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()
    stat = path.stat()
    meta: dict[str, Any] = {
        "filename": path.name,
        "filepath": str(path.resolve()),
        "file_type": ext.lstrip(".").upper(),
        "file_size": stat.st_size,
        "mtime": stat.st_mtime,
        "title": path.stem.replace("_", " ").title(),
        "char_count": 0,
        "page_count": 1,
        "text_content": "",
        "chapters": []
    }

    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            pages = []
            for i, page in enumerate(reader.pages):
                txt = (page.extract_text() or "").strip()
                if txt:
                    pages.append(f"--- Page {i + 1} ---\n{txt}")
            meta["page_count"] = len(reader.pages)
            meta["text_content"] = "\n\n".join(pages)
            meta["char_count"] = len(meta["text_content"])
            # PDF metadata title if present
            if reader.metadata and reader.metadata.title:
                meta["title"] = reader.metadata.title
        except Exception as e:
            meta["text_content"] = f"[Error reading PDF: {e}]"
        return meta

    if ext == ".epub":
        epub_data = parse_epub(path)
        meta["title"] = epub_data["title"]
        meta["author"] = epub_data.get("author", "Unknown")
        meta["chapters"] = epub_data["chapters"]
        meta["page_count"] = epub_data["page_count"]
        meta["text_content"] = epub_data["text_content"]
        meta["char_count"] = epub_data["char_count"]
        return meta

    if ext in (".docx", ".doc"):
        try:
            import docx
            doc = docx.Document(str(path))
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        paras.append(row_txt)
            meta["text_content"] = "\n\n".join(paras)
            meta["char_count"] = len(meta["text_content"])
            meta["page_count"] = max(1, math.ceil(len(paras) / 5))
        except Exception as e:
            meta["text_content"] = f"[Error reading Word document: {e}]"
        return meta

    # Plaintext, Markdown, HTML, JSON, CSV
    try:
        raw = path.read_bytes()
        try:
            txt = raw.decode("utf-8")
        except UnicodeDecodeError:
            txt = raw.decode("latin-1", errors="replace")

        if ext == ".html":
            txt = strip_html_tags(txt)

        meta["text_content"] = txt
        meta["char_count"] = len(txt)
        meta["page_count"] = max(1, math.ceil(len(txt) / 3000))

        # Look for markdown header # Title
        if ext in (".md", ".txt"):
            header_match = re.search(r"^#\s+(.+)$", txt, re.MULTILINE)
            if header_match:
                meta["title"] = header_match.group(1).strip()
    except Exception as e:
        meta["text_content"] = f"[Error reading text document: {e}]"

    return meta


# ==============================================================================
# Watch-Folder Manager & FTS5 Indexing
# ==============================================================================

def add_watch_folder(folder_path: str, name: str, workspace_id: str = "ws_default", recursive: bool = True) -> dict:
    """Add a local directory to be watched and indexed."""
    p = Path(folder_path).expanduser().resolve()
    if not p.is_dir():
        raise ValueError(f"Directory path does not exist: {folder_path}")

    source_id = f"src_{uuid.uuid4().hex[:10]}"
    with db.conn() as c:
        c.execute(
            """INSERT INTO knowledge_sources (id, workspace_id, folder_path, name, recursive, auto_sync)
               VALUES (?, ?, ?, ?, ?, 1)""",
            (source_id, workspace_id, str(p), name.strip() or p.name, 1 if recursive else 0)
        )
    # Perform immediate initial sync
    sync_watch_folder(source_id)
    return get_watch_folder(source_id) or {}


def get_watch_folder(source_id: str) -> dict | None:
    with db.conn() as c:
        row = c.execute("SELECT * FROM knowledge_sources WHERE id=?", (source_id,)).fetchone()
        return dict(row) if row else None


def list_watch_folders(workspace_id: str | None = None) -> list[dict]:
    with db.conn() as c:
        if workspace_id:
            rows = c.execute(
                "SELECT * FROM knowledge_sources WHERE workspace_id=? ORDER BY created_at DESC",
                (workspace_id,)
            ).fetchall()
        else:
            rows = c.execute("SELECT * FROM knowledge_sources ORDER BY created_at DESC").fetchall()
        return [dict(r) for r in rows]


def remove_watch_folder(source_id: str) -> bool:
    """Remove watched folder and prune its indexed documents."""
    with db.conn() as c:
        # Get documents to delete from FTS
        doc_rows = c.execute("SELECT id FROM knowledge_documents WHERE source_id=?", (source_id,)).fetchall()
        for r in doc_rows:
            c.execute("DELETE FROM knowledge_fts WHERE doc_id=?", (r["id"],))
        c.execute("DELETE FROM knowledge_documents WHERE source_id=?", (source_id,))
        c.execute("DELETE FROM knowledge_sources WHERE id=?", (source_id,))
        return True


def sync_watch_folder(source_id: str) -> dict[str, Any]:
    """
    Scans the watched folder, detects added/modified/deleted files,
    updates metadata, and keeps SQLite FTS5 index in sync.
    """
    source = get_watch_folder(source_id)
    if not source:
        raise ValueError(f"Watch source not found: {source_id}")

    folder = Path(source["folder_path"])
    if not folder.is_dir():
        return {"error": f"Folder does not exist: {folder}"}

    workspace_id = source["workspace_id"] or "ws_default"
    recursive = bool(source["recursive"])

    # Discover candidate files
    pattern = "**/*" if recursive else "*"
    found_files: list[Path] = []
    for item in folder.glob(pattern):
        if item.is_file() and item.suffix.lower() in SUPPORTED_EXTENSIONS:
            found_files.append(item)

    added_count = 0
    updated_count = 0
    removed_count = 0

    with db.conn() as c:
        # Load existing indexed documents for this source
        existing = {
            r["filepath"]: dict(r)
            for r in c.execute("SELECT * FROM knowledge_documents WHERE source_id=?", (source_id,)).fetchall()
        }

        current_paths = set()
        for fpath in found_files:
            str_path = str(fpath.resolve())
            current_paths.add(str_path)
            stat = fpath.stat()
            mtime = stat.st_mtime
            size = stat.st_size

            old_doc = existing.get(str_path)
            if not old_doc:
                # New file to index
                try:
                    parsed = parse_document(fpath)
                    doc_id = f"kdoc_{uuid.uuid4().hex[:12]}"
                    now = datetime.datetime.now().isoformat()
                    c.execute(
                        """INSERT INTO knowledge_documents
                           (id, source_id, workspace_id, filename, filepath, file_type, file_size, title, char_count, page_count, summary, mtime, created_at, updated_at)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (doc_id, source_id, workspace_id, parsed["filename"], str_path, parsed["file_type"],
                         size, parsed["title"], parsed["char_count"], parsed["page_count"], "", mtime, now, now)
                    )
                    # Index in FTS5
                    c.execute(
                        """INSERT INTO knowledge_fts (doc_id, workspace_id, title, filename, content)
                           VALUES (?, ?, ?, ?, ?)""",
                        (doc_id, workspace_id, parsed["title"], parsed["filename"], parsed["text_content"])
                    )
                    added_count += 1
                except Exception as e:
                    logger.warning("Failed to parse and index %s: %s", fpath, e)

            elif old_doc["mtime"] != mtime or old_doc["file_size"] != size:
                # File modified
                try:
                    parsed = parse_document(fpath)
                    doc_id = old_doc["id"]
                    now = datetime.datetime.now().isoformat()
                    c.execute(
                        """UPDATE knowledge_documents
                           SET filename=?, file_type=?, file_size=?, title=?, char_count=?, page_count=?, mtime=?, updated_at=?
                           WHERE id=?""",
                        (parsed["filename"], parsed["file_type"], size, parsed["title"],
                         parsed["char_count"], parsed["page_count"], mtime, now, doc_id)
                    )
                    # Update FTS5
                    c.execute("DELETE FROM knowledge_fts WHERE doc_id=?", (doc_id,))
                    c.execute(
                        """INSERT INTO knowledge_fts (doc_id, workspace_id, title, filename, content)
                           VALUES (?, ?, ?, ?, ?)""",
                        (doc_id, workspace_id, parsed["title"], parsed["filename"], parsed["text_content"])
                    )
                    updated_count += 1
                except Exception as e:
                    logger.warning("Failed to re-index %s: %s", fpath, e)

        # Detect deleted files
        for old_path, old_doc in existing.items():
            if old_path not in current_paths:
                c.execute("DELETE FROM knowledge_fts WHERE doc_id=?", (old_doc["id"],))
                c.execute("DELETE FROM knowledge_documents WHERE id=?", (old_doc["id"],))
                removed_count += 1

        # Update source metadata
        now_iso = datetime.datetime.now().isoformat()
        total_files = len(current_paths)
        c.execute(
            "UPDATE knowledge_sources SET last_scanned_at=?, file_count=? WHERE id=?",
            (now_iso, total_files, source_id)
        )

    return {
        "ok": True,
        "source_id": source_id,
        "total_files": len(found_files),
        "added": added_count,
        "updated": updated_count,
        "removed": removed_count
    }


def sync_all_watch_folders(workspace_id: str | None = None) -> list[dict]:
    """Sync all watch folders belonging to a workspace or globally."""
    sources = list_watch_folders(workspace_id)
    results = []
    for s in sources:
        try:
            results.append(sync_watch_folder(s["id"]))
        except Exception as e:
            results.append({"source_id": s["id"], "error": str(e)})
    return results


# ==============================================================================
# Full-Text Search (FTS5) with Snippets & Source Citation
# ==============================================================================

def search_knowledge_fts(query: str, workspace_id: str | None = None, limit: int = 25) -> list[dict]:
    """
    Query the knowledge base using SQLite FTS5.
    Returns matched documents with rank and highlighted contextual snippets.
    """
    clean_q = re.sub(r"[^\w\s\*\"]+", " ", query).strip()
    if not clean_q:
        return []

    # Format tokens with prefix wildcards if needed
    tokens = [f"{t}*" if not t.endswith("*") and not t.startswith('"') else t for t in clean_q.split()]
    fts_match_query = " ".join(tokens)

    with db.conn() as c:
        try:
            if workspace_id:
                sql = """
                    SELECT
                        k.id, k.title, k.filename, k.file_type, k.file_size, k.page_count,
                        k.workspace_id, k.filepath,
                        snippet(knowledge_fts, 4, '<mark>', '</mark>', ' … ', 24) AS snippet,
                        bm25(knowledge_fts) AS rank
                    FROM knowledge_fts f
                    JOIN knowledge_documents k ON k.id = f.doc_id
                    WHERE knowledge_fts MATCH ? AND (f.workspace_id = ? OR f.workspace_id = 'ws_default')
                    ORDER BY rank ASC
                    LIMIT ?
                """
                rows = c.execute(sql, (fts_match_query, workspace_id, limit)).fetchall()
            else:
                sql = """
                    SELECT
                        k.id, k.title, k.filename, k.file_type, k.file_size, k.page_count,
                        k.workspace_id, k.filepath,
                        snippet(knowledge_fts, 4, '<mark>', '</mark>', ' … ', 24) AS snippet,
                        bm25(knowledge_fts) AS rank
                    FROM knowledge_fts f
                    JOIN knowledge_documents k ON k.id = f.doc_id
                    WHERE knowledge_fts MATCH ?
                    ORDER BY rank ASC
                    LIMIT ?
                """
                rows = c.execute(sql, (fts_match_query, limit)).fetchall()

            return [dict(r) for r in rows]
        except Exception as e:
            logger.warning("FTS search error (%s): %s", fts_match_query, e)
            return []


def list_knowledge_documents(workspace_id: str | None = None, limit: int = 100) -> list[dict]:
    """List indexed knowledge documents."""
    with db.conn() as c:
        if workspace_id:
            rows = c.execute(
                """SELECT id, source_id, workspace_id, filename, file_type, file_size, title,
                          char_count, page_count, summary, updated_at
                   FROM knowledge_documents
                   WHERE workspace_id=? OR workspace_id='ws_default'
                   ORDER BY updated_at DESC LIMIT ?""",
                (workspace_id, limit)
            ).fetchall()
        else:
            rows = c.execute(
                """SELECT id, source_id, workspace_id, filename, file_type, file_size, title,
                          char_count, page_count, summary, updated_at
                   FROM knowledge_documents
                   ORDER BY updated_at DESC LIMIT ?""",
                (limit,)
            ).fetchall()
        return [dict(r) for r in rows]


# ==============================================================================
# Reading Desk & Reading Progress
# ==============================================================================

def get_document_for_reading(doc_id: str) -> dict[str, Any] | None:
    """Retrieve full text and chapter structure for the Reading Desk."""
    with db.conn() as c:
        row = c.execute("SELECT * FROM knowledge_documents WHERE id=?", (doc_id,)).fetchone()
        if not row:
            # Fallback check library_documents (from direct uploads)
            lib_row = c.execute("SELECT * FROM library_documents WHERE id=?", (doc_id,)).fetchone()
            if not lib_row:
                return None
            doc = dict(lib_row)
            doc["title"] = doc.get("filename", "Document").replace("_", " ").title()
            doc["filepath"] = ""
            text = doc.get("extracted_text", "")
            progress_row = c.execute("SELECT * FROM reading_progress WHERE doc_id=?", (doc_id,)).fetchone()
            return {
                "id": doc_id,
                "title": doc["title"],
                "filename": doc["filename"],
                "file_type": doc.get("file_type", "TXT"),
                "file_size": doc.get("file_size", 0),
                "page_count": doc.get("page_count", 1),
                "text_content": text,
                "chapters": [],
                "progress": dict(progress_row) if progress_row else {"progress_pct": 0.0, "bookmarks_json": "[]"}
            }

        doc = dict(row)
        filepath = Path(doc["filepath"])
        chapters = []
        text = ""

        if filepath.is_file():
            parsed = parse_document(filepath)
            chapters = parsed.get("chapters", [])
            text = parsed.get("text_content", "")
        else:
            # Fetch from FTS if file is not on disk
            fts_row = c.execute("SELECT content FROM knowledge_fts WHERE doc_id=?", (doc_id,)).fetchone()
            if fts_row:
                text = fts_row["content"]

        progress_row = c.execute("SELECT * FROM reading_progress WHERE doc_id=?", (doc_id,)).fetchone()
        return {
            "id": doc["id"],
            "title": doc.get("title") or doc["filename"],
            "filename": doc["filename"],
            "file_type": doc["file_type"],
            "file_size": doc["file_size"],
            "page_count": doc.get("page_count", 1),
            "text_content": text,
            "chapters": chapters,
            "progress": dict(progress_row) if progress_row else {"progress_pct": 0.0, "bookmarks_json": "[]"}
        }


def save_reading_progress(doc_id: str, workspace_id: str = "ws_default", progress_pct: float = 0.0,
                          last_position: str = "", bookmarks: list[dict] | None = None) -> bool:
    """Save user reading progress and bookmarks."""
    rec_id = f"prog_{uuid.uuid4().hex[:10]}"
    bookmarks_str = json.dumps(bookmarks or [])
    now = datetime.datetime.now().isoformat()
    with db.conn() as c:
        c.execute(
            """INSERT INTO reading_progress (id, doc_id, workspace_id, progress_pct, last_position, bookmarks_json, last_read_at)
               VALUES (?, ?, ?, ?, ?, ?, ?)
               ON CONFLICT(doc_id) DO UPDATE SET
                   progress_pct=excluded.progress_pct,
                   last_position=excluded.last_position,
                   bookmarks_json=excluded.bookmarks_json,
                   last_read_at=excluded.last_read_at""",
            (rec_id, doc_id, workspace_id, progress_pct, last_position, bookmarks_str, now)
        )
        return True


# ==============================================================================
# Research Notebook & Sourced Citations
# ==============================================================================

def add_research_note(workspace_id: str, doc_id: str | None, doc_title: str,
                      passage: str, note: str = "", tags: str = "", citation: str = "") -> dict:
    """Save an extracted passage with thoughts, tags, and citation."""
    note_id = f"note_{uuid.uuid4().hex[:10]}"
    with db.conn() as c:
        c.execute(
            """INSERT INTO research_notes (id, workspace_id, doc_id, doc_title, passage, note, tags, citation)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (note_id, workspace_id, doc_id, doc_title, passage.strip(), note.strip(), tags.strip(), citation.strip())
        )
    return get_research_note(note_id) or {}


def get_research_note(note_id: str) -> dict | None:
    with db.conn() as c:
        row = c.execute("SELECT * FROM research_notes WHERE id=?", (note_id,)).fetchone()
        return dict(row) if row else None


def list_research_notes(workspace_id: str | None = None, doc_id: str | None = None) -> list[dict]:
    with db.conn() as c:
        query = "SELECT * FROM research_notes WHERE 1=1"
        params: list[Any] = []
        if workspace_id:
            query += " AND (workspace_id=? OR workspace_id='ws_default')"
            params.append(workspace_id)
        if doc_id:
            query += " AND doc_id=?"
            params.append(doc_id)
        query += " ORDER BY created_at DESC"
        rows = c.execute(query, tuple(params)).fetchall()
        return [dict(r) for r in rows]


def delete_research_note(note_id: str) -> bool:
    with db.conn() as c:
        c.execute("DELETE FROM research_notes WHERE id=?", (note_id,))
        return True


def export_research_notes_markdown(workspace_id: str | None = None) -> str:
    """Export all research notes into formatted GitHub-flavored Markdown."""
    notes = list_research_notes(workspace_id)
    if not notes:
        return "# Research Notebook\n\n*No notes recorded yet.*"

    lines = ["# Research Notebook & Sourced Knowledge\n"]
    lines.append(f"*Generated on {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}*\n")

    current_doc = None
    for n in sorted(notes, key=lambda x: x.get("doc_title") or ""):
        doc_name = n.get("doc_title") or "General Notes"
        if doc_name != current_doc:
            current_doc = doc_name
            lines.append(f"\n## 📖 {current_doc}\n")

        citation = f" *({n['citation']})*" if n.get("citation") else ""
        lines.append(f"> \"{n['passage']}\"{citation}\n")
        if n.get("note"):
            lines.append(f"**Note:** {n['note']}\n")
        if n.get("tags"):
            tag_badges = " ".join([f"`#{t.strip()}`" for t in n["tags"].split(",") if t.strip()])
            lines.append(f"**Tags:** {tag_badges}\n")
        lines.append("---\n")

    return "\n".join(lines)


# ==============================================================================
# Learning Studio & Spaced Repetition (SuperMemo SM-2)
# ==============================================================================

def create_flashcard(deck_id: str, deck_title: str, workspace_id: str, front: str, back: str) -> dict:
    """Create a new study flashcard."""
    card_id = f"card_{uuid.uuid4().hex[:10]}"
    now = datetime.datetime.now().isoformat()
    with db.conn() as c:
        c.execute(
            """INSERT INTO study_flashcards
               (id, deck_id, deck_title, workspace_id, front, back, difficulty, interval_days, repetitions, next_review_at, created_at)
               VALUES (?, ?, ?, ?, ?, ?, 2.5, 1, 0, ?, ?)""",
            (card_id, deck_id, deck_title, workspace_id, front.strip(), back.strip(), now, now)
        )
    return get_flashcard(card_id) or {}


def get_flashcard(card_id: str) -> dict | None:
    with db.conn() as c:
        row = c.execute("SELECT * FROM study_flashcards WHERE id=?", (card_id,)).fetchone()
        return dict(row) if row else None


def list_flashcards(workspace_id: str | None = None, deck_id: str | None = None, due_only: bool = False) -> list[dict]:
    """Retrieve flashcards with optional filtering by deck and review due date."""
    now_iso = datetime.datetime.now().isoformat()
    with db.conn() as c:
        query = "SELECT * FROM study_flashcards WHERE 1=1"
        params: list[Any] = []
        if workspace_id:
            query += " AND (workspace_id=? OR workspace_id='ws_default')"
            params.append(workspace_id)
        if deck_id:
            query += " AND deck_id=?"
            params.append(deck_id)
        if due_only:
            query += " AND (next_review_at IS NULL OR next_review_at <= ?)"
            params.append(now_iso)
        query += " ORDER BY next_review_at ASC"
        rows = c.execute(query, tuple(params)).fetchall()
        return [dict(r) for r in rows]


def list_flashcard_decks(workspace_id: str | None = None) -> list[dict]:
    """List all flashcard decks with card counts and due counts."""
    now_iso = datetime.datetime.now().isoformat()
    with db.conn() as c:
        ws_clause = "WHERE (workspace_id=? OR workspace_id='ws_default')" if workspace_id else ""
        params = (workspace_id,) if workspace_id else ()
        sql = f"""
            SELECT
                deck_id, deck_title, workspace_id,
                COUNT(*) as total_cards,
                SUM(CASE WHEN next_review_at IS NULL OR next_review_at <= '{now_iso}' THEN 1 ELSE 0 END) as due_cards,
                MAX(last_reviewed_at) as last_reviewed_at
            FROM study_flashcards
            {ws_clause}
            GROUP BY deck_id
            ORDER BY deck_title ASC
        """
        rows = c.execute(sql, params).fetchall()
        return [dict(r) for r in rows]


def review_flashcard_sm2(card_id: str, rating: int) -> dict:
    """
    SuperMemo SM-2 Spaced Repetition Algorithm.
    Rating scale:
      1 = Again (Blackout / Total failure)
      2 = Hard (Correct with significant hesitation)
      3 = Good (Correct after hesitation)
      4 = Easy (Instant, perfect recall)
    """
    card = get_flashcard(card_id)
    if not card:
        raise ValueError("Flashcard not found")

    q = max(1, min(4, rating))
    # Map 1-4 scale to standard SM-2 0-5 scale
    sm2_score = {1: 1, 2: 3, 3: 4, 4: 5}[q]

    easiness = card["difficulty"]
    reps = card["repetitions"]
    interval = card["interval_days"]

    if sm2_score < 3:
        # Failed recall: reset reps and interval
        reps = 0
        interval = 1
    else:
        # Successful recall
        if reps == 0:
            interval = 1
        elif reps == 1:
            interval = 6
        else:
            interval = max(1, math.ceil(interval * easiness))
        reps += 1

    # Adjust easiness factor: EF' = EF + (0.1 - (5 - q) * (0.08 + (5 - q) * 0.02))
    easiness = easiness + (0.1 - (5 - sm2_score) * (0.08 + (5 - sm2_score) * 0.02))
    easiness = max(1.3, min(3.5, easiness))

    now = datetime.datetime.now()
    next_review = now + datetime.timedelta(days=interval)

    with db.conn() as c:
        c.execute(
            """UPDATE study_flashcards
               SET difficulty=?, interval_days=?, repetitions=?, next_review_at=?, last_reviewed_at=?
               WHERE id=?""",
            (round(easiness, 2), interval, reps, next_review.isoformat(), now.isoformat(), card_id)
        )

    return {
        "ok": True,
        "card_id": card_id,
        "repetitions": reps,
        "interval_days": interval,
        "next_review_at": next_review.isoformat(),
        "difficulty": round(easiness, 2)
    }


def delete_flashcard(card_id: str) -> bool:
    with db.conn() as c:
        c.execute("DELETE FROM study_flashcards WHERE id=?", (card_id,))
        return True


# Run table initialization on import
init_knowledge_tables()
